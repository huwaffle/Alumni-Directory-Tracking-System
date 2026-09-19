from models.employment_model import EmploymentModel

from datetime import datetime
from tkinter import filedialog
from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Border,
    Side,
    Alignment
)

from openpyxl.utils import get_column_letter

from models.user_model import UserModel


def export_alumni():

    user_model = UserModel()

    alumni = user_model.get_all_alumni()

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Alumni"

    thin = Side(
        style="thin",
        color="000000"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    # =====================================
    # Report Title
    # =====================================

    sheet["A1"] = "DE LA SALLE ARANETA UNIVERSITY"
    sheet["A2"] = "Alumni Tracking System"
    sheet["A3"] = "Alumni Report"

    sheet["A5"] = (
        "Generated: "
        + datetime.now().strftime("%B %d, %Y %I:%M %p")
    )

    headers = [
        "Student Number",
        "First Name",
        "Middle Name",
        "Last Name",
        "Sex",
        "Age",
        "Civil Status",
        "Email",
        "Address",
        "Course",
        "Graduation Year",
        "Latin Honors",
        "Username",
        "Status"
    ]

    # Header Row
    for column, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=7,
            column=column
        )

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            start_color="00474F",
            end_color="00474F"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = border

        

    # Data Rows
    row = 8

    for user in alumni:

        sheet.cell(row=row, column=1).value = user["student_number"]
        sheet.cell(row=row, column=2).value = user["first_name"]
        sheet.cell(row=row, column=3).value = user["middle_name"]
        sheet.cell(row=row, column=4).value = user["last_name"]
        sheet.cell(row=row, column=5).value = user["sex"]
        sheet.cell(row=row, column=6).value = user["age"]
        sheet.cell(row=row, column=7).value = user["civil_status"]
        sheet.cell(row=row, column=8).value = user["email"]
        sheet.cell(row=row, column=9).value = user["address"]
        sheet.cell(row=row, column=10).value = user["course"]
        sheet.cell(row=row, column=11).value = user["graduation_year"]
        sheet.cell(row=row, column=12).value = user["latin_honors"]
        sheet.cell(row=row, column=13).value = user["username"]
        sheet.cell(row=row, column=14).value = user["status"]

        for column in range(1, 15):

            sheet.cell(
                row=row,
                column=column
            ).border = border

        row += 1

    for column in sheet.columns:

        length = 0

        column_letter = get_column_letter(column[0].column)

        for cell in column:

            try:

                if len(str(cell.value)) > length:
                    length = len(str(cell.value))
            except:
                pass

        sheet.column_dimensions[column_letter].width = length + 3

    sheet.freeze_panes = "A8"
    sheet.auto_filter.ref = sheet.dimensions
    
    # =====================================
    # Employment Report Sheet
    # =====================================

    employment_sheet = workbook.create_sheet(
        title="Employment Report"
    )
    
    employment_model = EmploymentModel()

    conn = employment_model.connect()
    conn.row_factory = __import__("sqlite3").Row

    cursor = conn.cursor()

    cursor.execute("""

    SELECT

        users.student_number,
        users.first_name,
        users.last_name,

        employment.employment_status,
        employment.company,
        employment.job_title,
        employment.industry,
        employment.work_location,
        employment.salary_range

    FROM employment

    JOIN users

    ON employment.user_id = users.id

    ORDER BY users.last_name

    """)

    employment_records = cursor.fetchall()

    conn.close()

    employment_sheet["A1"] = "DE LA SALLE ARANETA UNIVERSITY"
    employment_sheet["A2"] = "Alumni Tracking System"
    employment_sheet["A3"] = "Employment Report"

    employment_sheet["A5"] = (
        "Generated: "
        + datetime.now().strftime("%B %d, %Y %I:%M %p")
    )

    headers = [

        "Student Number",
        "First Name",
        "Last Name",
        "Employment Status",
        "Company",
        "Job Title",
        "Industry",
        "Work Location",
        "Salary Range"

    ]

    for column, header in enumerate(headers, start=1):

        cell = employment_sheet.cell(
            row=9,
            column=column
        )

        cell.value = header
        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            start_color="00474F",
            end_color="00474F"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

        cell.border = border

    row = 10

    for record in employment_records:

        employment_sheet.cell(row=row, column=1).value = record["student_number"]
        employment_sheet.cell(row=row, column=2).value = record["first_name"]
        employment_sheet.cell(row=row, column=3).value = record["last_name"]
        employment_sheet.cell(row=row, column=4).value = record["employment_status"]
        employment_sheet.cell(row=row, column=5).value = record["company"]
        employment_sheet.cell(row=row, column=6).value = record["job_title"]
        employment_sheet.cell(row=row, column=7).value = record["industry"]
        employment_sheet.cell(row=row, column=8).value = record["work_location"]
        employment_sheet.cell(row=row, column=9).value = record["salary_range"]

        for col in range(1, 10):

            employment_sheet.cell(
                row=row,
                column=col
            ).border = border

        row += 1
    
    for column in employment_sheet.columns:

        length = 0

        column_letter = get_column_letter(column[0].column)

        for cell in column:

            try:

                if len(str(cell.value)) > length:
                    length = len(str(cell.value))
            except:
                pass

        employment_sheet.column_dimensions[column_letter].width = length + 3

    employment_sheet.freeze_panes = "A10"
    employment_sheet.auto_filter.ref = employment_sheet.dimensions
    
    filename = filedialog.asksaveasfilename(

        defaultextension=".xlsx",

        filetypes=[
            ("Excel Workbook", "*.xlsx")
        ],

        initialfile="Alumni_Report.xlsx"
    )

    if filename:

        workbook.save(filename)

        return True

    return False
